import pptx
import sys
sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('国产图片生成模型测试分析报告.pptx')

for slide_idx in [0, 1, 2, 3, 4, 5, 6]:
    slide = prs.slides[slide_idx]
    print(f'=== Slide {slide_idx+1} ===')
    for shape_idx, shape in enumerate(slide.shapes):
        print(f'  Shape[{shape_idx}]: shape_type={shape.shape_type}')
        if hasattr(shape, 'text'):
            txt = shape.text
            print(f'    Text: {repr(txt[:200])}')
        if hasattr(shape, 'has_table') and shape.has_table:
            t = shape.table
            print(f'    Table: {len(t.rows)}rows x {len(t.columns)}cols')
            for r, row in enumerate(t.rows):
                cells = [cell.text for cell in row.cells]
                print(f'      Row{r}: {cells}')
    print()
